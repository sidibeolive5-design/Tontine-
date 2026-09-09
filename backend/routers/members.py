"""Invitations + direct member enrolment (for tontines already running)."""

import secrets
from typing import Any, Optional

from fastapi import APIRouter, Depends, HTTPException, Response
from pydantic import BaseModel, EmailStr, Field

from lib.auth import ensure_permission, hash_password, require_staff, set_session
from lib.core import (
    assert_gerance_access,
    audit,
    enroll_member,
    get_tontine,
    new_id,
    notify,
    now_utc,
    parse_date,
)
from lib.db import db

router = APIRouter()


class InvitationInput(BaseModel):
    first_name: str = Field(min_length=1)
    last_name: str = Field(min_length=1)
    phone: str = ""
    email: EmailStr
    tontine_id: Optional[str] = None


class InvitationOut(BaseModel):
    id: str
    token: str
    invite_path: str
    first_name: str
    last_name: str
    phone: str
    email: str
    role: str
    gerance_id: str
    gerance_name: str
    tontine_id: Optional[str] = None
    tontine_name: Optional[str] = None
    status: str
    created_at: Any
    accepted_at: Optional[Any] = None


async def _invitation_out(i: dict[str, Any]) -> InvitationOut:
    g = await db.gerances.find_one({"id": i["gerance_id"]}, {"_id": 0})
    t = await db.tontines.find_one({"id": i.get("tontine_id")}, {"_id": 0}) if i.get("tontine_id") else None
    return InvitationOut(
        **i,
        invite_path=f"/invitation?token={i['token']}",
        gerance_name=g["name"] if g else "—",
        tontine_name=t["name"] if t else None,
    )


@router.post("/invitations", response_model=InvitationOut)
async def create_invitation(payload: InvitationInput, user: dict[str, Any] = Depends(require_staff)):
    ensure_permission(user, "invite_members")
    if await db.users.find_one({"email": payload.email.lower()}):
        raise HTTPException(status_code=409, detail="Un compte existe déjà avec cet email")
    pending = await db.invitations.find_one({"email": payload.email.lower(), "status": "sent"}, {"_id": 0})
    if pending:
        raise HTTPException(status_code=409, detail="Une invitation est déjà en attente pour cet email")
    gerance_id = user.get("gerance_id")
    if not gerance_id:
        raise HTTPException(status_code=400, detail="Aucune gérance rattachée à ce compte")
    if payload.tontine_id:
        tontine = await get_tontine(payload.tontine_id)
        assert_gerance_access(user, tontine["gerance_id"])
        if tontine["gerance_id"] != gerance_id:
            raise HTTPException(status_code=403, detail="Cette tontine appartient à une autre gérance")
    doc = {
        "id": new_id(),
        "token": secrets.token_urlsafe(24),
        "first_name": payload.first_name.strip(),
        "last_name": payload.last_name.strip(),
        "phone": payload.phone.strip(),
        "email": payload.email.lower(),
        "role": "member",
        "gerance_id": gerance_id,
        "tontine_id": payload.tontine_id,
        "invited_by": user["id"],
        "status": "sent",
        "created_at": now_utc(),
        "accepted_at": None,
    }
    await db.invitations.insert_one(doc)
    await audit(user, "invitation_created", "invitation", doc["id"], gerance_id, {"email": doc["email"]})
    doc.pop("_id", None)
    return await _invitation_out(doc)


@router.get("/invitations", response_model=list[InvitationOut])
async def list_invitations(user: dict[str, Any] = Depends(require_staff), gerance_id: Optional[str] = None):
    if user["role"] == "admin":
        query: dict[str, Any] = {"gerance_id": gerance_id} if gerance_id else {}
    else:
        query = {"gerance_id": user.get("gerance_id")}
    rows = await db.invitations.find(query, {"_id": 0}).sort("created_at", -1).to_list(500)
    return [await _invitation_out(i) for i in rows]


@router.post("/invitations/{invitation_id}/cancel", response_model=InvitationOut)
async def cancel_invitation(invitation_id: str, user: dict[str, Any] = Depends(require_staff)):
    ensure_permission(user, "invite_members")
    inv = await db.invitations.find_one({"id": invitation_id}, {"_id": 0})
    if not inv:
        raise HTTPException(status_code=404, detail="Invitation introuvable")
    assert_gerance_access(user, inv["gerance_id"])
    if inv["status"] != "sent":
        raise HTTPException(status_code=409, detail="Cette invitation n'est plus en attente")
    await db.invitations.update_one({"id": invitation_id}, {"$set": {"status": "cancelled"}})
    await audit(user, "invitation_cancelled", "invitation", invitation_id, inv["gerance_id"])
    fresh = await db.invitations.find_one({"id": invitation_id}, {"_id": 0})
    assert fresh is not None
    return await _invitation_out(fresh)


class PublicInvitation(BaseModel):
    first_name: str
    last_name: str
    email: str
    gerance_name: str
    tontine_name: Optional[str] = None
    status: str


@router.get("/invitations/token/{token}", response_model=PublicInvitation)
async def read_invitation(token: str):
    inv = await db.invitations.find_one({"token": token}, {"_id": 0})
    if not inv:
        raise HTTPException(status_code=404, detail="Invitation introuvable ou expirée")
    full = await _invitation_out(inv)
    return PublicInvitation(
        first_name=full.first_name,
        last_name=full.last_name,
        email=full.email,
        gerance_name=full.gerance_name,
        tontine_name=full.tontine_name,
        status=full.status,
    )


class AcceptInvitation(BaseModel):
    password: str = Field(min_length=6)
    phone: Optional[str] = None


class AcceptedUser(BaseModel):
    id: str
    first_name: str
    last_name: str
    phone: str
    email: str
    role: str
    status: str
    gerance_id: Optional[str] = None
    permissions: list[str] = []
    profile_complete: bool = False
    identity_status: str = "none"
    address: Optional[str] = None
    extra_info: Optional[str] = None


@router.post("/invitations/token/{token}/accept", response_model=AcceptedUser)
async def accept_invitation(token: str, payload: AcceptInvitation, response: Response):
    inv = await db.invitations.find_one({"token": token}, {"_id": 0})
    if not inv:
        raise HTTPException(status_code=404, detail="Invitation introuvable ou expirée")
    if inv["status"] != "sent":
        raise HTTPException(status_code=409, detail="Cette invitation a déjà été utilisée ou annulée")
    if await db.users.find_one({"email": inv["email"]}):
        raise HTTPException(status_code=409, detail="Un compte existe déjà avec cet email")
    user = {
        "id": new_id(),
        "first_name": inv["first_name"],
        "last_name": inv["last_name"],
        "phone": (payload.phone or inv.get("phone") or "").strip(),
        "email": inv["email"],
        "password_hash": hash_password(payload.password),
        "role": "member",
        "status": "active",
        "gerance_id": None,
        "permissions": [],
        "profile_complete": False,
        "identity_status": "none",
        "address": None,
        "extra_info": None,
        "created_at": now_utc(),
    }
    await db.users.insert_one(user)
    await db.invitations.update_one(
        {"token": token}, {"$set": {"status": "accepted", "accepted_at": now_utc(), "user_id": user["id"]}}
    )
    if inv.get("tontine_id"):
        tontine = await get_tontine(inv["tontine_id"])
        await enroll_member(tontine, user["id"])
        await notify(
            user["id"],
            "Vous êtes inscrit à une tontine",
            f"Vous avez été ajouté à {tontine['name']}. Votre contrat est prêt à être signé.",
            tontine["gerance_id"],
            tontine["id"],
            "member_enrolled",
        )
    await notify(
        user["id"],
        "Bienvenue sur AIDONS-NOUS VIVANTS",
        "Votre compte a été créé via une invitation. Complétez vos informations personnelles.",
        inv["gerance_id"],
        event="account_created",
    )
    set_session(response, user["id"])
    return AcceptedUser(**{k: v for k, v in user.items() if k in AcceptedUser.model_fields})


class DirectMemberInput(BaseModel):
    first_name: str = Field(min_length=1)
    last_name: str = Field(min_length=1)
    phone: str = ""
    email: EmailStr
    password: str = Field(min_length=6)
    tontine_id: Optional[str] = None


@router.post("/members/create")
async def create_member(payload: DirectMemberInput, user: dict[str, Any] = Depends(require_staff)):
    ensure_permission(user, "manage_members")
    if await db.users.find_one({"email": payload.email.lower()}):
        raise HTTPException(status_code=409, detail="Un compte existe déjà avec cet email")
    if payload.tontine_id:
        tontine = await get_tontine(payload.tontine_id)
        assert_gerance_access(user, tontine["gerance_id"])
    member = {
        "id": new_id(),
        "first_name": payload.first_name.strip(),
        "last_name": payload.last_name.strip(),
        "phone": payload.phone.strip(),
        "email": payload.email.lower(),
        "password_hash": hash_password(payload.password),
        "role": "member",
        "status": "active",
        "gerance_id": None,
        "permissions": [],
        "profile_complete": False,
        "identity_status": "none",
        "address": None,
        "extra_info": None,
        "created_at": now_utc(),
    }
    await db.users.insert_one(member)
    enrolled_in = None
    if payload.tontine_id:
        tontine = await get_tontine(payload.tontine_id)
        await enroll_member(tontine, member["id"])
        enrolled_in = tontine["name"]
        await notify(
            member["id"],
            "Vous êtes inscrit à une tontine",
            f"Vous avez été ajouté à {tontine['name']} par le responsable de la gérance.",
            tontine["gerance_id"],
            tontine["id"],
            "member_enrolled",
        )
    await audit(user, "member_created", "user", member["id"], user.get("gerance_id"), {"email": member["email"]})
    return {
        "id": member["id"],
        "first_name": member["first_name"],
        "last_name": member["last_name"],
        "email": member["email"],
        "enrolled_in": enrolled_in,
    }


class EnrolExisting(BaseModel):
    member_id: str
    tontine_id: str


@router.post("/members/enrol")
async def enrol_existing(payload: EnrolExisting, user: dict[str, Any] = Depends(require_staff)):
    ensure_permission(user, "manage_members")
    tontine = await get_tontine(payload.tontine_id)
    assert_gerance_access(user, tontine["gerance_id"])
    member = await db.users.find_one({"id": payload.member_id, "role": "member"}, {"_id": 0})
    if not member:
        raise HTTPException(status_code=404, detail="Membre introuvable")
    if await db.tontine_members.find_one({"tontine_id": tontine["id"], "member_id": member["id"]}):
        raise HTTPException(status_code=409, detail="Ce membre participe déjà à cette tontine")
    await enroll_member(tontine, member["id"])
    await audit(user, "member_enrolled", "tontine_member", member["id"], tontine["gerance_id"],
                {"tontine_id": tontine["id"]})
    await notify(
        member["id"],
        "Vous êtes inscrit à une tontine",
        f"Vous avez été ajouté à {tontine['name']} par le responsable de la gérance.",
        tontine["gerance_id"],
        tontine["id"],
        "member_enrolled",
    )
    return {"ok": True, "tontine_name": tontine["name"]}


class ImportInput(BaseModel):
    tontine_id: str
    file_base64: str = Field(min_length=10)  # data URL or raw base64 of a .xlsx / .csv
    filename: str = "membres.xlsx"


class ImportRow(BaseModel):
    line: int
    first_name: str
    last_name: str
    email: str
    status: str  # created | enrolled | skipped
    message: str


class ImportResult(BaseModel):
    tontine_name: str
    created: int
    enrolled: int
    skipped: int
    rows: list[ImportRow]


HEADER_ALIASES = {
    "prenom": "first_name", "prénom": "first_name", "first_name": "first_name", "firstname": "first_name",
    "nom": "last_name", "last_name": "last_name", "lastname": "last_name",
    "email": "email", "mail": "email", "e-mail": "email",
    "telephone": "phone", "téléphone": "phone", "phone": "phone", "tel": "phone",
}


def _parse_sheet(raw: bytes, filename: str) -> list[dict[str, str]]:
    """Read a .xlsx (openpyxl) or .csv into normalised dict rows."""
    import csv
    import io

    if filename.lower().endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:2000]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        table = [[(c or "").strip() for c in row] for row in reader]
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return []
        table = [[("" if c is None else str(c)).strip() for c in row] for row in ws.iter_rows(values_only=True)]
    table = [r for r in table if any(r)]
    if not table:
        return []
    header = [HEADER_ALIASES.get(h.strip().lower(), "") for h in table[0]]
    if "email" not in header:
        raise HTTPException(
            status_code=422,
            detail="Colonnes attendues : Prénom, Nom, Email, Téléphone (la colonne Email est obligatoire)",
        )
    rows: list[dict[str, str]] = []
    for raw_row in table[1:]:
        row = {"first_name": "", "last_name": "", "email": "", "phone": ""}
        for key, value in zip(header, raw_row):
            if key:
                row[key] = value
        rows.append(row)
    return rows


@router.post("/members/import", response_model=ImportResult)
async def import_members(payload: ImportInput, user: dict[str, Any] = Depends(require_staff)):
    """Bulk-enrol the members of a tontine already running, from an Excel/CSV sheet."""
    ensure_permission(user, "manage_members")
    tontine = await get_tontine(payload.tontine_id)
    assert_gerance_access(user, tontine["gerance_id"])
    import base64
    import binascii

    encoded = payload.file_base64.split(",", 1)[-1]
    try:
        raw = base64.b64decode(encoded, validate=False)
    except (binascii.Error, ValueError):
        raise HTTPException(status_code=422, detail="Fichier illisible")
    try:
        sheet = _parse_sheet(raw, payload.filename)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="Fichier illisible : utilisez un .xlsx ou un .csv")
    if not sheet:
        raise HTTPException(status_code=422, detail="Le fichier ne contient aucune ligne exploitable")

    out: list[ImportRow] = []
    created = enrolled = skipped = 0
    for index, row in enumerate(sheet, start=2):
        email = row["email"].strip().lower()
        first = row["first_name"].strip()
        last = row["last_name"].strip()
        if not email or "@" not in email:
            skipped += 1
            out.append(ImportRow(line=index, first_name=first, last_name=last, email=row["email"],
                                 status="skipped", message="Email manquant ou invalide"))
            continue
        existing = await db.users.find_one({"email": email}, {"_id": 0})
        if existing:
            if existing["role"] != "member":
                skipped += 1
                out.append(ImportRow(line=index, first_name=first, last_name=last, email=email,
                                     status="skipped", message="Ce compte n'est pas un compte membre"))
                continue
            member_id = existing["id"]
            already = await db.tontine_members.find_one({"tontine_id": tontine["id"], "member_id": member_id})
            if already:
                skipped += 1
                out.append(ImportRow(line=index, first_name=existing["first_name"], last_name=existing["last_name"],
                                     email=email, status="skipped", message="Participe déjà à cette tontine"))
                continue
            await enroll_member(tontine, member_id)
            enrolled += 1
            out.append(ImportRow(line=index, first_name=existing["first_name"], last_name=existing["last_name"],
                                 email=email, status="enrolled", message="Compte existant rattaché à la tontine"))
        else:
            if not first or not last:
                skipped += 1
                out.append(ImportRow(line=index, first_name=first, last_name=last, email=email,
                                     status="skipped", message="Prénom et nom obligatoires pour un nouveau compte"))
                continue
            member_id = new_id()
            await db.users.insert_one({
                "id": member_id, "first_name": first, "last_name": last, "phone": row["phone"].strip(),
                "email": email, "password_hash": hash_password(secrets.token_urlsafe(12)), "role": "member",
                "status": "active", "gerance_id": None, "permissions": [], "profile_complete": False,
                "identity_status": "none", "address": None, "extra_info": None, "created_at": now_utc(),
            })
            await enroll_member(tontine, member_id)
            invitation = {
                "id": new_id(), "token": secrets.token_urlsafe(24), "first_name": first, "last_name": last,
                "phone": row["phone"].strip(), "email": email, "role": "member",
                "gerance_id": tontine["gerance_id"], "tontine_id": tontine["id"], "invited_by": user["id"],
                "status": "cancelled", "created_at": now_utc(), "accepted_at": None,
                "note": "compte créé par import — mot de passe à définir par le responsable",
            }
            await db.invitations.insert_one(invitation)
            created += 1
            out.append(ImportRow(line=index, first_name=first, last_name=last, email=email,
                                 status="created", message="Compte créé et rattaché à la tontine"))
        await notify(
            member_id,
            "Vous êtes inscrit à une tontine",
            f"Vous avez été ajouté à {tontine['name']} par le responsable de la gérance.",
            tontine["gerance_id"],
            tontine["id"],
            "member_enrolled",
        )
    await audit(user, "members_imported", "tontine", tontine["id"], tontine["gerance_id"],
                {"created": created, "enrolled": enrolled, "skipped": skipped, "filename": payload.filename})
    return ImportResult(tontine_name=tontine["name"], created=created, enrolled=enrolled, skipped=skipped, rows=out)


class BulkRegularise(BaseModel):
    tontine_id: str
    member_id: str
    up_to_date: str  # ISO date: every due date <= this day is marked
    status: str = "paid"


@router.post("/due-dates/bulk-history")
async def bulk_history(payload: BulkRegularise, user: dict[str, Any] = Depends(require_staff)):
    """Historical catch-up for a tontine already running — never presented as a Wave payment."""
    ensure_permission(user, "record_history")
    if payload.status not in ("paid", "pending"):
        raise HTTPException(status_code=422, detail="Statut invalide")
    tontine = await get_tontine(payload.tontine_id)
    assert_gerance_access(user, tontine["gerance_id"])
    try:
        parse_date(payload.up_to_date)
    except ValueError:
        raise HTTPException(status_code=422, detail="Date invalide (format AAAA-MM-JJ)")
    label = "administrateur" if user["role"] == "admin" else "gérant"
    res = await db.contribution_due_dates.update_many(
        {
            "tontine_id": payload.tontine_id,
            "member_id": payload.member_id,
            "date": {"$lte": payload.up_to_date},
            "status": {"$ne": "processing"},
        },
        {"$set": {"status": payload.status, "source": f"historique_{label}", "synced_at": now_utc()}},
    )
    await audit(user, "due_dates_bulk_history", "tontine", payload.tontine_id, tontine["gerance_id"],
                {"member_id": payload.member_id, "up_to_date": payload.up_to_date, "status": payload.status,
                 "updated": res.modified_count})
    await notify(
        payload.member_id,
        "Historique enregistré",
        f"Historique enregistré par l'{label} : {res.modified_count} jour(s) jusqu'au {payload.up_to_date} "
        f"marqué(s) comme {'payé' if payload.status == 'paid' else 'impayé'} pour {tontine['name']}.",
        tontine["gerance_id"],
        tontine["id"],
        "history_recorded",
    )
    return {"ok": True, "updated": res.modified_count, "source": f"historique_{label}"}
